import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx_report(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles setup
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    i = 0
    in_code_block = False
    code_lines = []
    code_lang = ""

    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        # Code block delimiter
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lines = []
                code_lang = line[3:].strip()
            else:
                in_code_block = False
                # Render code block in table or shaded box
                if code_lang.lower() == 'mermaid':
                    # Add description for diagram
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(6)
                    run = p.add_run("[Diagramma di Flusso dell'Architettura - Vedere diagramma Mermaid nel report Markdown]")
                    run.italic = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(100, 116, 139)
                else:
                    tbl = doc.add_table(rows=1, cols=1)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    cell = tbl.cell(0, 0)
                    set_cell_background(cell, "F1F5F9")
                    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    run = p.add_run("\n".join(code_lines))
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(15, 23, 42)
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(2)
                p_spacer.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Blank lines
        if not line.strip():
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[2:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Deep navy
            i += 1
            continue

        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138) # Navy blue
            i += 1
            continue

        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(line[4:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.bold = True
            run.font.color.rgb = RGBColor(14, 116, 144) # Cyan/slate
            i += 1
            continue

        # Markdown Table Detection
        if '|' in line and i + 1 < len(lines) and '|' in lines[i+1] and ('---' in lines[i+1] or ':---' in lines[i+1]):
            table_lines = []
            while i < len(lines) and '|' in lines[i] and lines[i].strip():
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 3:
                # Parse headers
                raw_headers = [c.strip() for c in table_lines[0].strip('|').split('|')]
                num_cols = len(raw_headers)
                
                # Parse data rows
                data_rows = []
                for r_line in table_lines[2:]:
                    cols = [c.strip() for c in r_line.strip('|').split('|')]
                    if len(cols) < num_cols:
                        cols += [""] * (num_cols - len(cols))
                    elif len(cols) > num_cols:
                        cols = cols[:num_cols]
                    data_rows.append(cols)

                doc_tbl = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
                doc_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

                # Format header
                hdr_cells = doc_tbl.rows[0].cells
                for idx, h_text in enumerate(raw_headers):
                    cell = hdr_cells[idx]
                    set_cell_background(cell, "1E3A8A") # Navy header
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    clean_h = re.sub(r'\*\*(.*?)\*\*', r'\1', h_text)
                    run = p.add_run(clean_h)
                    run.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(255, 255, 255)

                # Format data rows
                for r_idx, r_cols in enumerate(data_rows):
                    row_cells = doc_tbl.rows[r_idx + 1].cells
                    bg_col = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
                    for c_idx, val in enumerate(r_cols):
                        cell = row_cells[c_idx]
                        set_cell_background(cell, bg_col)
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        
                        # Parse simple bold/italic in table cells
                        sub_parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', val)
                        for sp in sub_parts:
                            if sp.startswith('**') and sp.endswith('**'):
                                r = p.add_run(sp[2:-2])
                                r.bold = True
                            elif sp.startswith('*') and sp.endswith('*'):
                                r = p.add_run(sp[1:-1])
                                r.italic = True
                            else:
                                r = p.add_run(sp)
                            r.font.size = Pt(9.0)
                            r.font.color.rgb = RGBColor(30, 41, 59)

                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(4)
                p_spacer.paragraph_format.space_after = Pt(4)
            continue

        # Bullet lists
        if line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            content = line[2:].strip()
            _format_inlines(p, content)
            i += 1
            continue

        # Numbered lists
        m_num = re.match(r'^(\d+)\.\s+(.*)', line)
        if m_num:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            content = m_num.group(2).strip()
            _format_inlines(p, content)
            i += 1
            continue

        # Horizontal rule
        if line.startswith('---'):
            i += 1
            continue

        # Normal paragraphs
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        _format_inlines(p, line.strip())
        i += 1

    doc.save(docx_path)
    print(f"Documento DOCX generato con successo: {docx_path}")

def _format_inlines(paragraph, text):
    """Parses markdown bold, italic, and inline code spans."""
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 83, 9) # Amber/brown code
        elif token.startswith('$') and token.endswith('$'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Cambria Math'
            run.italic = True
        else:
            paragraph.add_run(token)

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(base_dir, "REPORT_TECNICO_PIPELINE_ECDC.md")
    out_docx = os.path.join(base_dir, "REPORT_TECNICO_PIPELINE_ECDC.docx")
    reports_docx = os.path.join(base_dir, "output", "reports", "REPORT_TECNICO_PIPELINE_ECDC.docx")

    create_docx_report(md_file, out_docx)
    create_docx_report(md_file, reports_docx)
