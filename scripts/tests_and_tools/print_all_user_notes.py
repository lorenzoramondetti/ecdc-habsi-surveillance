import docx
import re

doc = docx.Document('REPORT_TECNICO_PIPELINE_ECDC Con correzioni.docx')

print("=== DETAILED REVIEW OF ALL BRACKETED USER NOTES AND EDITS ===")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    # Check for [ ... ]
    brackets = re.findall(r'\[(.*?)\]', text)
    has_note = False
    for b in brackets:
        # filter out simple confidence intervals or table notations
        if not re.match(r'^\d+(\.\d+)?%?\s*[-–—]\s*\d+(\.\d+)?%?$', b.strip()):
            has_note = True
    if has_note or '?' in text or 'VORREI' in text or 'INSERIRE' in text or 'CONTROLLARE' in text:
        print(f"\n--- [P{i+1}] ---")
        print(text)

print("\n=== CHECKING ALL TABLES FOR NOTES ===")
for ti, t in enumerate(doc.tables):
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            ctext = cell.text.strip()
            if any(k in ctext for k in ['[', '?', 'VORREI', 'INSERIRE', 'CONTROLLARE', 'controllare']):
                if not re.search(r'\[\d+.*%\]', ctext):
                    print(f"Table {ti+1} Row {ri+1} Col {ci+1}: {ctext}")
