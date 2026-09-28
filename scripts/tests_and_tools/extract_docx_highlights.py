import docx
from docx.oxml.ns import qn

doc = docx.Document('REPORT_TECNICO_PIPELINE_ECDC Con correzioni.docx')

print("Checking paragraphs and runs XML tags...")
found_any = False

for i, p in enumerate(doc.paragraphs):
    p_xml = p._p.xml
    # check for highlight, shd, color, comment, ins, del
    for tag in ['highlight', 'shd', 'color', 'comment', 'ins', 'del']:
        if tag in p_xml:
            # print run details
            for r in p.runs:
                r_xml = r._r.xml
                if any(t in r_xml for t in ['highlight', 'shd', 'color', 'ins', 'del']):
                    print(f"P{i+1}: text='{r.text}' | XML={r_xml[:200]}")
                    found_any = True
                    break

for ti, t in enumerate(doc.tables):
    t_xml = t._tbl.xml
    for tag in ['highlight', 'shd', 'color', 'comment', 'ins', 'del']:
        if tag in t_xml:
            for ri, row in enumerate(t.rows):
                for ci, cell in enumerate(row.cells):
                    cell_xml = cell._tc.xml
                    if any(s in cell_xml for s in ['highlight', 'shd', 'color', 'comment']):
                        for p in cell.paragraphs:
                            for r in p.runs:
                                if any(s in r._r.xml for s in ['highlight', 'shd', 'color']):
                                    print(f"T{ti+1} R{ri+1} C{ci+1}: text='{r.text}' | XML={r._r.xml[:200]}")
                                    found_any = True

print(f"Done. Found any: {found_any}")
