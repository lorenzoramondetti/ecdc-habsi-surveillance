import docx
import difflib
import re

doc_corr = docx.Document('REPORT_TECNICO_PIPELINE_ECDC Con correzioni.docx')
doc_orig = docx.Document('REPORT_TECNICO_PIPELINE_ECDC.docx')

c_paras = [p.text.strip() for p in doc_corr.paragraphs if p.text.strip()]
o_paras = [p.text.strip() for p in doc_orig.paragraphs if p.text.strip()]

print("=== ALL BRACKETED NOTES IN CORREZIONI ===")
for i, p in enumerate(doc_corr.paragraphs):
    matches = re.findall(r'\[.*?\]', p.text)
    for m in matches:
        if not re.match(r'^\[\d+.*%\]$', m) and not re.match(r'^\[\d+\]$', m): # ignore CI intervals like [51.01% - 100.0%]
            print(f"P{i+1}: {p.text}")
            break

print("\n=== COMPLETE DIFF OF PARAGRAPHS ===")
matcher = difflib.SequenceMatcher(None, o_paras, c_paras)
diff_num = 1
for tag, i1, i2, j1, j2 in matcher.get_opcodes():
    if tag != 'equal':
        print(f"\n--- CHANGE #{diff_num} ({tag}) ---")
        if i1 < i2:
            print("ORIGINAL:")
            for k in range(i1, i2):
                print(f"  {o_paras[k]}")
        if j1 < j2:
            print("CORREZIONI:")
            for k in range(j1, j2):
                print(f"  {c_paras[k]}")
        diff_num += 1

print("\n=== CHECKING TABLES ===")
for ti, (t_orig, t_corr) in enumerate(zip(doc_orig.tables, doc_corr.tables)):
    orig_rows = [[c.text.strip().replace('\n', ' ') for c in r.cells] for r in t_orig.rows]
    corr_rows = [[c.text.strip().replace('\n', ' ') for c in r.cells] for r in t_corr.rows]
    if orig_rows != corr_rows:
        print(f"\nTable {ti+1} differences:")
        for r_idx, (ro, rc) in enumerate(zip(orig_rows, corr_rows)):
            if ro != rc:
                print(f"  Row {r_idx+1}:")
                print(f"    ORIG: {ro}")
                print(f"    CORR: {rc}")
