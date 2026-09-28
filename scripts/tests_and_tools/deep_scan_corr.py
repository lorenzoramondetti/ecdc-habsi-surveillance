import docx
import difflib

doc_corr = docx.Document('REPORT_TECNICO_PIPELINE_ECDC Con correzioni.docx')
doc_orig = docx.Document('REPORT_TECNICO_PIPELINE_ECDC.docx')

print("--- DIFFING PARAGRAPHS BETWEEN ORIGINAL AND CORREZIONI ---")
corr_paras = [p.text for p in doc_corr.paragraphs]
orig_paras = [p.text for p in doc_orig.paragraphs]

diff = list(difflib.unified_diff(orig_paras, corr_paras, lineterm='', fromfile='orig', tofile='corr'))
if diff:
    print(f"Differences found in text ({len(diff)} diff lines):")
    for d in diff[:50]:
        print(d)
else:
    print("No text differences between paragraphs!")

print("\n--- CHECKING RUN-LEVEL PROPERTIES IN CORREZIONI ---")
for i, p in enumerate(doc_corr.paragraphs):
    for r in p.runs:
        rPr = r._r.rPr
        if rPr is not None:
            # check all child tags in rPr
            children = [c.tag.split('}')[-1] for c in rPr]
            # check for anything interesting
            interesting = [c for c in rPr if any(k in c.tag for k in ['highlight', 'shd', 'color', 'strike', 'u', 'b', 'i', 'vanish'])]
            # print if highlight or shd or color is not standard
            for c in rPr:
                tag = c.tag.split('}')[-1]
                val = c.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '')
                fill = c.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill', '')
                if tag in ['highlight', 'shd'] or (tag == 'color' and val.lower() not in ['1f4e79', '2e75b6', '333333', '595959', '000000', 'auto', '']):
                    print(f"P{i+1} [{tag} val={val} fill={fill}]: '{r.text}'")

for ti, t in enumerate(doc_corr.tables):
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            tcPr = cell._tc.tcPr
            if tcPr is not None:
                for c in tcPr:
                    tag = c.tag.split('}')[-1]
                    fill = c.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill', '')
                    # check if fill is yellow or similar
                    if fill and fill.lower() not in ['1f4e79', 'f2f4f7', 'ffffff', 'auto', 'none']:
                        print(f"Table {ti+1} R{ri+1} C{ci+1} Cell [{tag} fill={fill}]: {[p.text for p in cell.paragraphs]}")
            for p in cell.paragraphs:
                for r in p.runs:
                    rPr = r._r.rPr
                    if rPr is not None:
                        for c in rPr:
                            tag = c.tag.split('}')[-1]
                            val = c.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '')
                            fill = c.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill', '')
                            if tag in ['highlight', 'shd'] or (tag == 'color' and val.lower() not in ['1f4e79', '2e75b6', '333333', '595959', '000000', 'auto', '']):
                                print(f"Table {ti+1} R{ri+1} C{ci+1} [{tag} val={val} fill={fill}]: '{r.text}'")
